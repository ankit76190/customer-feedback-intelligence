import io
import csv
import openpyxl
from app.models.feedback import FeedbackRecord
from app.models.analysis import FeedbackItemAnalysis, SentimentEnum, UrgencyLevel
from app.services.analytics_engine import AnalyticsEngine
from app.services.excel_exporter import ExcelExporterService, sanitize_formula_injection

def test_excel_export_structure_and_sheet_names():
    records = [
        FeedbackRecord(
            id="FB-1",
            original_row_number=2,
            original_text="Very fast interface",
            cleaned_text="Very fast interface",
            rating=5.0,
            category="In-App",
            analysis=FeedbackItemAnalysis(
                sentiment=SentimentEnum.POSITIVE,
                sentiment_reason="Customer praised the interface speed",
                primary_theme="Performance",
                secondary_themes=["Speed"],
                pain_points=[],
                positive_aspects=["Fast interface"],
                urgency=UrgencyLevel.LOW,
                summary="Fast interface praise."
            )
        )
    ]
    insights = AnalyticsEngine.calculate_insights(records)
    stream = ExcelExporterService.generate_workbook(records, insights)
    
    assert stream is not None
    assert stream.getbuffer().nbytes > 1000

    wb = openpyxl.load_workbook(stream)
    sheet_names = wb.sheetnames
    assert "Feedback" in sheet_names
    assert "Sentiment Summary" in sheet_names
    assert "Theme Summary" in sheet_names
    assert "Data Quality" in sheet_names

    # Check Feedback sheet has 2 rows (header + 1 record)
    ws_feedback = wb["Feedback"]
    assert ws_feedback.max_row == 2
    assert ws_feedback["A2"].value == "FB-1"
    assert ws_feedback["H2"].value == "Very fast interface"

def test_spreadsheet_formula_injection_defense():
    """Verify that malicious formula prefixes (=, +, -, @) are sanitized with a leading quote."""
    assert sanitize_formula_injection("=SUM(A1:A10)") == "'=SUM(A1:A10)"
    assert sanitize_formula_injection("+cmd|' /C calc'!A0") == "'+cmd|' /C calc'!A0"
    assert sanitize_formula_injection("-5+2") == "'-5+2"
    assert sanitize_formula_injection("@SUM(1,2)") == "'@SUM(1,2)"
    assert sanitize_formula_injection("Normal feedback text") == "Normal feedback text"
    assert sanitize_formula_injection(42) == 42
    assert sanitize_formula_injection(None) == ""

    # Test in actual workbook export
    unsafe_records = [
        FeedbackRecord(
            id="=MALICIOUS-ID",
            original_row_number=2,
            original_text="=cmd|'/C calc'!A0",
            cleaned_text="-dangerous_calc",
            rating=1.0,
            analysis=FeedbackItemAnalysis(
                sentiment=SentimentEnum.NEGATIVE,
                sentiment_reason="@alert('hacked')",
                primary_theme="+ThemeInjection",
                urgency=UrgencyLevel.HIGH,
                summary="=HYPERLINK('http://evil.com')"
            )
        )
    ]
    insights = AnalyticsEngine.calculate_insights(unsafe_records)
    stream = ExcelExporterService.generate_workbook(unsafe_records, insights)
    wb = openpyxl.load_workbook(stream)
    ws = wb["Feedback"]

    # In Excel, a cell value starting with '= is treated as literal text
    assert ws["A2"].value.startswith("'")
    assert ws["H2"].value.startswith("'")
    assert ws["I2"].value.startswith("'")

def test_csv_export_format_and_unicode():
    records = [
        FeedbackRecord(
            id="FB-UNICODE",
            original_row_number=1,
            original_text="Très bon produit! 🚀 Fast & reliable.",
            cleaned_text="Très bon produit! 🚀 Fast & reliable.",
            rating=5.0,
            category="In-App",
            analysis=FeedbackItemAnalysis(
                sentiment=SentimentEnum.POSITIVE,
                sentiment_reason="Customer expressed high satisfaction",
                primary_theme="Usability",
                urgency=UrgencyLevel.LOW,
                summary="High satisfaction expressed."
            )
        )
    ]
    csv_stream = ExcelExporterService.generate_csv(records)
    assert csv_stream is not None
    
    # Check UTF-8 with BOM bytes
    raw_bytes = csv_stream.getvalue()
    assert raw_bytes.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM
    
    text = raw_bytes.decode("utf-8-sig")
    assert "Très bon produit! 🚀" in text
    assert "FB-UNICODE" in text
    assert "Positive" in text
