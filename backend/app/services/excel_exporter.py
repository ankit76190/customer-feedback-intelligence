import io
import csv
from datetime import datetime
from typing import List, Optional, Any
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from app.models.feedback import FeedbackRecord
from app.models.analysis import ExecutiveInsights

def sanitize_formula_injection(val: Any) -> Any:
    """
    Prevents spreadsheet formula injection (CSV/Excel formula injection).
    If a text value starts with '=', '+', '-', '@', '\t', or '\r',
    it is prepended with a single quote so Excel/Calc treats it strictly as passive text.
    """
    if val is None:
        return ""
    if isinstance(val, (int, float, bool)):
        return val
    s = str(val)
    if s and s[0] in ("=", "+", "-", "@", "\t", "\r"):
        return f"'{s}"
    return s

class ExcelExporterService:
    """
    Generates enterprise-grade, multi-sheet Excel workbooks and sanitized CSV exports.
    Adheres strictly to Phase 5 requirements:
    - 4 distinct sheets: 'Feedback', 'Sentiment Summary', 'Theme Summary', and 'Data Quality'.
    - Safe formula injection prevention across all verbatim and text cells.
    - Readable column widths, frozen headers, auto-filters, and soft badge styling.
    - Never includes API keys or internal stack traces.
    """

    COLOR_PRIMARY = "1E293B"      # Slate 800
    COLOR_HEADER_BG = "0F172A"    # Slate 900
    COLOR_WHITE = "FFFFFF"
    COLOR_ZEBRA = "F8FAFC"        # Slate 50
    COLOR_BORDER = "CBD5E1"       # Slate 300
    
    COLOR_POS_BG = "D1FAE5"       # Emerald 100
    COLOR_POS_TXT = "065F46"      # Emerald 800
    COLOR_NEG_BG = "FFE4E6"       # Rose 100
    COLOR_NEG_TXT = "9F1239"      # Rose 800
    COLOR_NEU_BG = "FEF3C7"       # Amber 100
    COLOR_NEU_TXT = "92400E"      # Amber 800
    COLOR_REV_BG = "F3E8FF"       # Purple 100
    COLOR_REV_TXT = "6B21A8"      # Purple 800

    @classmethod
    def generate_workbook(cls, records: List[FeedbackRecord], insights: ExecutiveInsights) -> io.BytesIO:
        wb = openpyxl.Workbook()
        
        # Sheet 1: Feedback (Full Verbatim Ledger with Analysis)
        ws_feedback = wb.active
        ws_feedback.title = "Feedback"
        cls._build_feedback_sheet(ws_feedback, records)

        # Sheet 2: Sentiment Summary (Executive overview, health, and distribution)
        ws_sentiment = wb.create_sheet(title="Sentiment Summary")
        cls._build_sentiment_summary_sheet(ws_sentiment, insights)

        # Sheet 3: Theme Summary (Recurring themes, pain points, and opportunities)
        ws_themes = wb.create_sheet(title="Theme Summary")
        cls._build_theme_summary_sheet(ws_themes, insights)

        # Sheet 4: Data Quality (Cleaning audit, flags, and validity breakdown)
        ws_quality = wb.create_sheet(title="Data Quality")
        cls._build_data_quality_sheet(ws_quality, records)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @classmethod
    def _build_feedback_sheet(cls, ws, records: List[FeedbackRecord]):
        """Sheet 1: Feedback — Full ledger with analysis status, verbatims, themes, and pain points."""
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A2"

        headers = [
            "Record ID", "Row #", "Date", "Customer ID", "Product", "Channel / Category", "Rating",
            "Original Verbatim", "Cleaned Text", "Sentiment", "Confidence", 
            "Primary Theme", "Secondary Themes", "Pain Points", "Summary", "Analysis Status"
        ]

        # Style header row
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_idx, value=header)
            cell.font = Font(name="Calibri", size=10, bold=True, color=cls.COLOR_WHITE)
            cell.fill = PatternFill(start_color=cls.COLOR_HEADER_BG, end_color=cls.COLOR_HEADER_BG, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[1].height = 28

        thin_border = Border(
            left=Side(style='thin', color=cls.COLOR_BORDER),
            right=Side(style='thin', color=cls.COLOR_BORDER),
            top=Side(style='thin', color=cls.COLOR_BORDER),
            bottom=Side(style='thin', color=cls.COLOR_BORDER)
        )

        for row_idx, r in enumerate(records, 2):
            analysis = r.analysis
            if analysis:
                sent_str = analysis.sentiment.value if hasattr(analysis.sentiment, "value") else str(analysis.sentiment)
                status_str = "Needs Review" if "needs_review" in sent_str.lower() else "Analyzed"
                sentiment = sent_str.capitalize().replace("_", " ")
                conf = f"{round(analysis.confidence * 100)}%" if analysis.confidence is not None else ""
                theme = analysis.primary_theme or ""
                sec_themes = ", ".join(analysis.secondary_themes) if analysis.secondary_themes else ""
                pain_pts = "; ".join(analysis.pain_points) if analysis.pain_points else ""
                summary = analysis.summary or ""
            else:
                sentiment = "Unanalyzed"
                status_str = "Pending"
                conf = ""
                theme = ""
                sec_themes = ""
                pain_pts = ""
                summary = ""

            row_data = [
                sanitize_formula_injection(r.id),
                r.original_row_number,
                sanitize_formula_injection(r.date or ""),
                sanitize_formula_injection(r.customer_id or ""),
                sanitize_formula_injection(r.product or ""),
                sanitize_formula_injection(r.category or ""),
                r.rating if r.rating is not None else "",
                sanitize_formula_injection(r.original_text),
                sanitize_formula_injection(r.cleaned_text),
                sentiment,
                conf,
                sanitize_formula_injection(theme),
                sanitize_formula_injection(sec_themes),
                sanitize_formula_injection(pain_pts),
                sanitize_formula_injection(summary),
                status_str
            ]

            is_zebra = (row_idx % 2 == 0)
            row_fill = PatternFill(start_color=cls.COLOR_ZEBRA, end_color=cls.COLOR_ZEBRA, fill_type="solid") if is_zebra else None

            for col_idx, val in enumerate(row_data, 1):
                cell = ws.cell(row=row_idx, column=col_idx, value=val)
                cell.font = Font(name="Calibri", size=9)
                cell.border = thin_border
                if row_fill:
                    cell.fill = row_fill

                # Alignment
                if col_idx in [1, 2, 3, 4, 7, 10, 11, 16]:
                    cell.alignment = Alignment(horizontal="center", vertical="top")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)

                # Sentiment conditional styling (col 10)
                if col_idx == 10:
                    s_low = str(sentiment).lower()
                    if "pos" in s_low:
                        cell.fill = PatternFill(start_color=cls.COLOR_POS_BG, end_color=cls.COLOR_POS_BG, fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color=cls.COLOR_POS_TXT)
                    elif "neg" in s_low:
                        cell.fill = PatternFill(start_color=cls.COLOR_NEG_BG, end_color=cls.COLOR_NEG_BG, fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color=cls.COLOR_NEG_TXT)
                    elif "neu" in s_low:
                        cell.fill = PatternFill(start_color=cls.COLOR_NEU_BG, end_color=cls.COLOR_NEU_BG, fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color=cls.COLOR_NEU_TXT)
                    elif "review" in s_low:
                        cell.fill = PatternFill(start_color=cls.COLOR_REV_BG, end_color=cls.COLOR_REV_BG, fill_type="solid")
                        cell.font = Font(name="Calibri", size=9, bold=True, color=cls.COLOR_REV_TXT)

            ws.row_dimensions[row_idx].height = 24

        # Enable auto-filter across all columns
        max_col_letter = get_column_letter(len(headers))
        ws.auto_filter.ref = f"A1:{max_col_letter}{max(2, len(records) + 1)}"

        # Readable column widths
        widths = {
            1: 14, 2: 8, 3: 16, 4: 15, 5: 18, 6: 18, 7: 10,
            8: 48, 9: 48, 10: 14, 11: 12, 12: 20, 13: 24, 14: 32,
            15: 36, 16: 16
        }
        for col_idx, width in widths.items():
            ws.column_dimensions[get_column_letter(col_idx)].width = width

    @classmethod
    def _build_sentiment_summary_sheet(cls, ws, insights: ExecutiveInsights):
        """Sheet 2: Sentiment Summary — Executive KPIs, Net Sentiment, and sentiment breakdown."""
        ws.views.sheetView[0].showGridLines = True

        # Header Title
        ws.merge_cells("A1:G1")
        title_cell = ws["A1"]
        title_cell.value = "Customer Feedback Intelligence — Sentiment & Executive Summary"
        title_cell.font = Font(name="Calibri", size=15, bold=True, color=cls.COLOR_WHITE)
        title_cell.fill = PatternFill(start_color=cls.COLOR_HEADER_BG, end_color=cls.COLOR_HEADER_BG, fill_type="solid")
        title_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[1].height = 36

        ws["A2"].value = f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')} | Analysis Engine: Google Gemini 2.5"
        ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="64748B")

        # KPI Block
        kpis = [
            ("Total Records", insights.total_records),
            ("Classified Records", insights.classified_records or (insights.total_records - insights.unclassified_records)),
            ("Completion Rate", f"{insights.completion_rate}%"),
            ("Net Sentiment Score", f"{insights.net_sentiment_score:+0.1f}")
        ]

        row = 4
        for idx, (label, val) in enumerate(kpis):
            col_letter = get_column_letter(idx * 2 + 1)
            next_col = get_column_letter(idx * 2 + 2)
            ws.merge_cells(f"{col_letter}{row}:{next_col}{row}")
            ws.merge_cells(f"{col_letter}{row+1}:{next_col}{row+1}")
            
            lbl_cell = ws[f"{col_letter}{row}"]
            lbl_cell.value = label.upper()
            lbl_cell.font = Font(name="Calibri", size=9, bold=True, color="64748B")
            lbl_cell.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
            lbl_cell.alignment = Alignment(horizontal="center", vertical="center")

            val_cell = ws[f"{col_letter}{row+1}"]
            val_cell.value = str(val)
            val_cell.font = Font(name="Calibri", size=13, bold=True, color=cls.COLOR_PRIMARY)
            val_cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
            val_cell.alignment = Alignment(horizontal="center", vertical="center")

        ws.row_dimensions[4].height = 20
        ws.row_dimensions[5].height = 28

        # Executive Summary Box
        ws.merge_cells("A7:H7")
        ws["A7"].value = "EXECUTIVE AI INSIGHTS SYNTHESIS"
        ws["A7"].font = Font(name="Calibri", size=10, bold=True, color=cls.COLOR_WHITE)
        ws["A7"].fill = PatternFill(start_color=cls.COLOR_PRIMARY, end_color=cls.COLOR_PRIMARY, fill_type="solid")
        ws["A7"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

        ws.merge_cells("A8:H10")
        summary_cell = ws["A8"]
        summary_cell.value = sanitize_formula_injection(insights.ai_executive_summary)
        summary_cell.font = Font(name="Calibri", size=10, color="1E293B")
        summary_cell.alignment = Alignment(wrap_text=True, vertical="top")
        summary_cell.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

        # Sentiment Distribution Table
        ws["A12"].value = "SENTIMENT DISTRIBUTION & DENOMINATORS"
        ws["A12"].font = Font(name="Calibri", size=11, bold=True, color=cls.COLOR_PRIMARY)

        headers_sent = ["Sentiment Classification", "Observed Count", "Distribution (%)", "Denominator Base", "Status & Notes"]
        for col_idx, h in enumerate(headers_sent, 1):
            cell = ws.cell(row=13, column=col_idx, value=h)
            cell.font = Font(name="Calibri", size=10, bold=True, color=cls.COLOR_WHITE)
            cell.fill = PatternFill(start_color=cls.COLOR_PRIMARY, end_color=cls.COLOR_PRIMARY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        classified_denom = insights.classified_records or (insights.total_records - insights.unclassified_records)
        s_rows = [
            ("Positive", insights.sentiment_counts.get("Positive", 0), f"{insights.sentiment_percentages.get('Positive', 0.0)}%", f"{classified_denom} classified", "Customer compliments and strengths", cls.COLOR_POS_BG, cls.COLOR_POS_TXT),
            ("Neutral", insights.sentiment_counts.get("Neutral", 0), f"{insights.sentiment_percentages.get('Neutral', 0.0)}%", f"{classified_denom} classified", "Balanced inquiries and operational feedback", cls.COLOR_NEU_BG, cls.COLOR_NEU_TXT),
            ("Negative", insights.sentiment_counts.get("Negative", 0), f"{insights.sentiment_percentages.get('Negative', 0.0)}%", f"{classified_denom} classified", "Customer friction points and complaints", cls.COLOR_NEG_BG, cls.COLOR_NEG_TXT),
            ("Needs Review", insights.needs_review_count or insights.sentiment_counts.get("Needs_Review", 0), "Excluded", f"{insights.total_records} total", "Flagged for human manual verification", cls.COLOR_REV_BG, cls.COLOR_REV_TXT),
        ]

        curr_r = 14
        for label, cnt, pct, denom, notes, bg, txt in s_rows:
            ws.cell(row=curr_r, column=1, value=label).font = Font(name="Calibri", size=10, bold=True, color=txt)
            ws.cell(row=curr_r, column=1).fill = PatternFill(start_color=bg, end_color=bg, fill_type="solid")
            ws.cell(row=curr_r, column=2, value=cnt).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=3, value=pct).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=4, value=denom).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=5, value=notes).alignment = Alignment(horizontal="left")
            curr_r += 1

        # Denominator Disclosures
        curr_r += 1
        ws.merge_cells(f"A{curr_r}:H{curr_r}")
        disc_cell = ws[f"A{curr_r}"]
        disc_cell.value = f"Data Transparency Note: {insights.exclusions_disclaimer}"
        disc_cell.font = Font(name="Calibri", size=9, italic=True, color="64748B")

        for col_idx in range(1, 9):
            ws.column_dimensions[get_column_letter(col_idx)].width = 20

    @classmethod
    def _build_theme_summary_sheet(cls, ws, insights: ExecutiveInsights):
        """Sheet 3: Theme Summary — Thematic volume, recurring pain points, and opportunities."""
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A3"

        # Section 1: Themes Table
        ws["A1"].value = "RECURRING CUSTOMER THEMES"
        ws["A1"].font = Font(name="Calibri", size=12, bold=True, color=cls.COLOR_HEADER_BG)
        
        headers = ["Theme Name", "Observed Volume", "Share (%)", "Positive", "Neutral", "Negative"]
        for idx, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=idx, value=h)
            cell.font = Font(name="Calibri", size=10, bold=True, color=cls.COLOR_WHITE)
            cell.fill = PatternFill(start_color=cls.COLOR_PRIMARY, end_color=cls.COLOR_PRIMARY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_r = 3
        for tm in insights.top_recurring_themes:
            ws.cell(row=curr_r, column=1, value=sanitize_formula_injection(tm.theme)).font = Font(name="Calibri", size=9, bold=True)
            ws.cell(row=curr_r, column=2, value=tm.count).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=3, value=f"{tm.percentage}%").alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=4, value=tm.positive_count).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=5, value=tm.neutral_count).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=6, value=tm.negative_count).alignment = Alignment(horizontal="center")
            curr_r += 1

        # Section 2: Pain Points Table
        curr_r += 2
        ws.cell(row=curr_r, column=1, value="CRITICAL CUSTOMER PAIN POINTS").font = Font(name="Calibri", size=12, bold=True, color=cls.COLOR_HEADER_BG)
        curr_r += 1

        pp_headers = ["Pain Point", "Frequency", "Urgency Rating", "Representative Customer Verbatim"]
        for idx, h in enumerate(pp_headers, 1):
            cell = ws.cell(row=curr_r, column=idx, value=h)
            cell.font = Font(name="Calibri", size=10, bold=True, color=cls.COLOR_WHITE)
            cell.fill = PatternFill(start_color=cls.COLOR_PRIMARY, end_color=cls.COLOR_PRIMARY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_r += 1
        for pp in insights.critical_pain_points:
            ws.cell(row=curr_r, column=1, value=sanitize_formula_injection(pp.pain_point)).font = Font(name="Calibri", size=9, bold=True)
            ws.cell(row=curr_r, column=2, value=pp.frequency).alignment = Alignment(horizontal="center")
            
            urg_val = pp.urgency.value if hasattr(pp.urgency, "value") else str(pp.urgency)
            urg_cell = ws.cell(row=curr_r, column=3, value=urg_val)
            urg_cell.alignment = Alignment(horizontal="center")
            if urg_val.lower() == "high":
                urg_cell.fill = PatternFill(start_color=cls.COLOR_NEG_BG, end_color=cls.COLOR_NEG_BG, fill_type="solid")
                urg_cell.font = Font(name="Calibri", size=9, bold=True, color=cls.COLOR_NEG_TXT)
            elif urg_val.lower() == "medium":
                urg_cell.fill = PatternFill(start_color=cls.COLOR_NEU_BG, end_color=cls.COLOR_NEU_BG, fill_type="solid")
                urg_cell.font = Font(name="Calibri", size=9, bold=True, color=cls.COLOR_NEU_TXT)

            quote_cell = ws.cell(row=curr_r, column=4, value=sanitize_formula_injection(pp.sample_quote or ""))
            quote_cell.alignment = Alignment(wrap_text=True, vertical="top")
            curr_r += 1

        # Section 3: Actionable Improvement Opportunities
        curr_r += 2
        ws.cell(row=curr_r, column=1, value="RECOMMENDED ACTIONABLE OPPORTUNITIES").font = Font(name="Calibri", size=12, bold=True, color=cls.COLOR_HEADER_BG)
        curr_r += 1

        for idx, opp in enumerate(insights.top_improvement_opportunities, 1):
            ws.cell(row=curr_r, column=1, value=f"{idx}.").alignment = Alignment(horizontal="center")
            ws.merge_cells(f"B{curr_r}:D{curr_r}")
            opp_cell = ws.cell(row=curr_r, column=2, value=sanitize_formula_injection(opp))
            opp_cell.font = Font(name="Calibri", size=9)
            curr_r += 1

        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 14
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 55
        ws.column_dimensions["E"].width = 12
        ws.column_dimensions["F"].width = 12

    @classmethod
    def _build_data_quality_sheet(cls, ws, records: List[FeedbackRecord]):
        """Sheet 4: Data Quality — Validation status, normalization flags, and audit log."""
        ws.views.sheetView[0].showGridLines = True
        ws.freeze_panes = "A3"

        ws["A1"].value = "DATA QUALITY & NORMALIZATION AUDIT LOG"
        ws["A1"].font = Font(name="Calibri", size=12, bold=True, color=cls.COLOR_HEADER_BG)

        headers = ["Record ID", "Original Row #", "Validation Status", "Date Validity", "Audit / Flagged Reasons", "Cleaned Text Preview"]
        for idx, h in enumerate(headers, 1):
            cell = ws.cell(row=2, column=idx, value=h)
            cell.font = Font(name="Calibri", size=10, bold=True, color=cls.COLOR_WHITE)
            cell.fill = PatternFill(start_color=cls.COLOR_PRIMARY, end_color=cls.COLOR_PRIMARY, fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        curr_r = 3
        for r in records:
            ws.cell(row=curr_r, column=1, value=sanitize_formula_injection(r.id)).alignment = Alignment(horizontal="center")
            ws.cell(row=curr_r, column=2, value=r.original_row_number).alignment = Alignment(horizontal="center")
            
            status_cell = ws.cell(row=curr_r, column=3, value="Valid" if r.is_valid else "Flagged")
            status_cell.alignment = Alignment(horizontal="center")
            if r.is_valid:
                status_cell.font = Font(name="Calibri", size=9, color=cls.COLOR_POS_TXT, bold=True)
            else:
                status_cell.font = Font(name="Calibri", size=9, color=cls.COLOR_NEG_TXT, bold=True)
                status_cell.fill = PatternFill(start_color=cls.COLOR_NEG_BG, end_color=cls.COLOR_NEG_BG, fill_type="solid")

            date_status = "Valid Timestamp" if r.date else "Missing / Undated"
            ws.cell(row=curr_r, column=4, value=date_status).alignment = Alignment(horizontal="center")

            flags_str = ", ".join(r.flagged_reasons) if r.flagged_reasons else "None (Passed all cleaning rules)"
            ws.cell(row=curr_r, column=5, value=sanitize_formula_injection(flags_str))
            
            cleaned_preview = (r.cleaned_text[:80] + "...") if len(r.cleaned_text) > 80 else r.cleaned_text
            ws.cell(row=curr_r, column=6, value=sanitize_formula_injection(cleaned_preview))
            curr_r += 1

        ws.column_dimensions["A"].width = 16
        ws.column_dimensions["B"].width = 12
        ws.column_dimensions["C"].width = 16
        ws.column_dimensions["D"].width = 18
        ws.column_dimensions["E"].width = 38
        ws.column_dimensions["F"].width = 50

    @classmethod
    def generate_csv(cls, records: List[FeedbackRecord]) -> io.BytesIO:
        """
        Generates a clean, formula-injection-safe CSV export encoded with UTF-8 with BOM (utf-8-sig)
        for seamless cross-platform Excel and text editor compatibility.
        """
        output = io.StringIO()
        writer = csv.writer(output, quoting=csv.QUOTE_MINIMAL)

        headers = [
            "Record ID", "Row #", "Date", "Customer ID", "Product", "Channel / Category", "Rating",
            "Original Verbatim", "Cleaned Text", "Sentiment", "Confidence",
            "Primary Theme", "Secondary Themes", "Pain Points", "Summary", "Analysis Status"
        ]
        writer.writerow(headers)

        for r in records:
            analysis = r.analysis
            if analysis:
                sent_str = analysis.sentiment.value if hasattr(analysis.sentiment, "value") else str(analysis.sentiment)
                status_str = "Needs Review" if "needs_review" in sent_str.lower() else "Analyzed"
                sentiment = sent_str.capitalize().replace("_", " ")
                conf = f"{round(analysis.confidence * 100)}%" if analysis.confidence is not None else ""
                theme = analysis.primary_theme or ""
                sec_themes = ", ".join(analysis.secondary_themes) if analysis.secondary_themes else ""
                pain_pts = "; ".join(analysis.pain_points) if analysis.pain_points else ""
                summary = analysis.summary or ""
            else:
                sentiment = "Unanalyzed"
                status_str = "Pending"
                conf = ""
                theme = ""
                sec_themes = ""
                pain_pts = ""
                summary = ""

            writer.writerow([
                sanitize_formula_injection(r.id),
                r.original_row_number,
                sanitize_formula_injection(r.date or ""),
                sanitize_formula_injection(r.customer_id or ""),
                sanitize_formula_injection(r.product or ""),
                sanitize_formula_injection(r.category or ""),
                r.rating if r.rating is not None else "",
                sanitize_formula_injection(r.original_text),
                sanitize_formula_injection(r.cleaned_text),
                sentiment,
                conf,
                sanitize_formula_injection(theme),
                sanitize_formula_injection(sec_themes),
                sanitize_formula_injection(pain_pts),
                sanitize_formula_injection(summary),
                status_str
            ])

        byte_output = io.BytesIO(output.getvalue().encode("utf-8-sig"))
        byte_output.seek(0)
        return byte_output
